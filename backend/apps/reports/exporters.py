import csv
from decimal import Decimal
from io import BytesIO, StringIO

from django.http import HttpResponse
from django.utils import timezone

CREDIT = "Glideinbir — Built by Sahil Thakur"


def _cell(value):
    if isinstance(value, Decimal):
        return float(value)
    return value


def _filename(report, ext):
    return f"{report['name']}_{report['period']['from']}_{report['period']['to']}.{ext}"


def to_csv(report):
    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow([report["title"], f"{report['period']['from']} to {report['period']['to']}"])
    writer.writerow(report["columns"])
    for row in report["rows"]:
        # Prefix formula-like values so spreadsheets never execute them (CSV injection).
        writer.writerow([("'" + str(v)) if str(v)[:1] in ("=", "+", "-", "@") and not isinstance(v, (int, float,
                         Decimal)) else v for v in (row.get(c) for c in report["columns"])])
    writer.writerow([])
    writer.writerow([CREDIT, f"Generated {timezone.localtime():%Y-%m-%d %H:%M}"])
    response = HttpResponse(buf.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "csv")}"'
    return response


def to_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = report["name"][:31]
    ws.append([report["title"]])
    ws["A1"].font = Font(bold=True, size=13, color="4535C9")
    ws.append([f"Period: {report['period']['from']} to {report['period']['to']}"])
    ws.append([])
    ws.append([c.replace("_", " ").title() for c in report["columns"]])
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4535C9")
    for row in report["rows"]:
        ws.append([_cell(row.get(c)) for c in report["columns"]])
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 48)
    ws.append([])
    ws.append([CREDIT])
    wb.properties.creator = "Sahil Thakur"
    out = BytesIO()
    wb.save(out)
    response = HttpResponse(out.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "xlsx")}"'
    return response
